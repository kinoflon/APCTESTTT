from docx import Document
from docx.shared import Inches
import csv
from io import StringIO
import logging

def export_document(file_path, requester_info, comparison_result, worker_manager=None, alternatives_text=None):
    """Export the data to a formatted Word document file and save chat log.
    
    Args:
        file_path: Path to save the document
        requester_info: Dictionary containing requester information
        comparison_result: String containing the comparison result
        worker_manager: WorkerManager instance for exporting chat log
        alternatives_text: String containing the suggested alternatives text
    """
    doc = Document()
    
    # Title
    title = doc.add_heading("EOL Part Replacement Suggestion Form", level=1)
    title.alignment = 1  # Center alignment

    # Function to add centered section headers
    def add_section_header(text):
        header = doc.add_heading(text, level=2)
        header.alignment = 1
        
    # Function to add horizontal line
    def add_horizontal_line():
        paragraph = doc.add_paragraph()
        paragraph.add_run('_' * 100)
        paragraph.alignment = 1
        doc.add_paragraph()  # Add space after line

    # Section 1: Requester Information
    add_section_header("Section 1: Requester Information")
    for field in ["Name", "Department", "Position", "Date"]:
        value = requester_info.get(field, "")
        doc.add_paragraph(f"{field}: {value}")
    add_horizontal_line()

    # Section 2: EOL Part Details
    add_section_header("Section 2: EOL Part Details")
    eol_fields = [
        "EOL Part Name", "EOL Part Number", "Description", "Work Order",
        "Reason for Replacement"
    ]
    for field in eol_fields:
        value = requester_info.get(field, "")
        if field in ["Description", "Work Order"]:
            doc.add_paragraph(f"{field}:")
            doc.add_paragraph(value, style='List Bullet')
        else:
            doc.add_paragraph(f"{field}: {value}")
    add_horizontal_line()

    # Section 3: Comparison Table
    add_section_header("Section 3: Comparison Table")
    parts = comparison_result.split('\n\n', 2)
    if len(parts) >= 2:
        csv_table = parts[0]
        try:
            csv_reader = csv.reader(StringIO(csv_table))
            table_data = list(csv_reader)

            if table_data:
                # Transpose the table data
                transposed_data = list(zip(*table_data))
                
                # Create the table with transposed dimensions
                table = doc.add_table(rows=len(transposed_data), cols=len(transposed_data[0]))
                table.style = 'Table Grid'

                for i, row in enumerate(transposed_data):
                    for j, cell_value in enumerate(row):
                        cell = table.cell(i, j)
                        cell.text = cell_value.strip()
                        if j == 0:  # Bold the first column (specs)
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.bold = True
                table.autofit = True
            else:
                doc.add_paragraph("No data available for comparison table.")
        except Exception as e:
            doc.add_paragraph(f"Error creating table: {str(e)}")
    else:
        doc.add_paragraph("No comparison table data available.")
    add_horizontal_line()

    # Section 4: Comparison Summary
    add_section_header("Section 4: Comparison Summary")
    if len(parts) >= 2:
        analysis_text = '\n\n'.join(parts[1:])  # Join all remaining parts
        paragraphs = analysis_text.split('\n')
        for paragraph in paragraphs:
            if paragraph.strip():  # Only add non-empty paragraphs
                doc.add_paragraph(paragraph.strip())
    else:
        doc.add_paragraph("No analysis and recommendation available.")
    add_horizontal_line()
    
    # Section 5: Suggested Alternatives
    add_section_header("Section 5: Suggested Alternatives")
    if alternatives_text and alternatives_text.strip():
        # Add the alternatives text
        paragraphs = alternatives_text.split('\n')
        for paragraph in paragraphs:
            if paragraph.strip():  # Only add non-empty paragraphs
                doc.add_paragraph(paragraph.strip())
    else:
        doc.add_paragraph("No suggested alternatives available.")
    add_horizontal_line()

    # Section 6: Additional Comments
    add_section_header("Section 6: Additional Comments")
    doc.add_paragraph("[User can add comments here after export]")
    for _ in range(5):
        doc.add_paragraph("_" * 50)
    add_horizontal_line()

    # Section 7: Signatures
    add_section_header("Section 7: Signatures")
    doc.add_paragraph("Procurement Manager:")
    doc.add_paragraph("Name: ____________________")
    doc.add_paragraph("Signature: ____________________")
    doc.add_paragraph("Date: ____________________")
    doc.add_paragraph()
    doc.add_paragraph("Engineering Reviewer:")
    doc.add_paragraph("Name: ____________________")
    doc.add_paragraph("Signature: ____________________")
    doc.add_paragraph("Date: ____________________")

    # Save the document
    doc.save(file_path)
    
    # Save chat log if worker_manager is provided
    if worker_manager:
        # Create chat log filename by replacing .docx with .txt
        chat_log_path = file_path.rsplit('.', 1)[0] + '_chat_log.txt'
        if worker_manager.export_chat_log(chat_log_path):
            logging.info(f"Chat log saved to: {chat_log_path}")
        else:
            logging.error("Failed to save chat log")
